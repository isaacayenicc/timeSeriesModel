"""Artifact integrity checks and source/results release packaging."""
import hashlib
import json
import re
import zipfile
from pathlib import Path
import nbformat
import pandas as pd
from pptx import Presentation
from .io_matlab import sha256


def verify(root):
    """Fail if required deliverables, predictions, deck structures or links break."""
    root=Path(root)
    paths=['README.md','reports/PROJECT_FINAL_REPORT.md','docs/COMPREHENSIVE_STUDY_COMPANION.md',
           'notebooks/00_COMPLETE_PROJECT_WALKTHROUGH.ipynb','presentations/Progress_Presentation_8min.pptx',
           'presentations/Final_Presentation_12min.pptx','reports/results_summary.csv',
           'models/checkpoints_manifest.json','research/literature_matrix.csv']
    for name in paths:
        if not (root/name).is_file() or (root/name).stat().st_size==0:
            raise AssertionError('Missing required artifact: '+name)
    df=pd.read_csv(root/'reports/results_summary.csv')
    assert df.n_targets.nunique()==1 and (df.start_index==256).all() and (df.end_index==23999).all()
    for row in json.loads((root/'models/checkpoints_manifest.json').read_text()):
        matches=list((root/'models').glob(row['run_id']+'.*'))
        assert len(matches)==1 and sha256(matches[0])==row['artifact_sha256']
    fault_summary_path=root/'reports/fault_study_summary.json'
    fault_status='not executed'
    if fault_summary_path.exists():
        fault=json.loads(fault_summary_path.read_text())
        if fault.get('status')=='complete_descriptive_unlabeled':
            scalar=pd.read_csv(root/'reports/fault_record_features.csv')
            angular=pd.read_csv(root/'reports/fault_angular_features.csv')
            models={row['run_id'] for row in fault['selected_models']}|{'raw'}
            assert fault['n_records']==200 and fault['n_models']==13
            assert len(scalar)==200*len(models) and len(angular)==3*len(scalar)
            assert set(scalar.model_id)==models and set(angular.model_id)==models
            assert all(set(g.record_id)==set(range(1,201)) for _,g in scalar.groupby('model_id'))
            assert all(set(g.record_id)==set(range(1,201)) for _,g in angular.groupby(['model_id','speed_unit_assumption']))
            assert (scalar.first_target_index==256).all() and (scalar.last_target_index==31999).all()
            assert (scalar.n_evaluated==31744).all()
            assert set(angular.speed_unit_assumption)=={'rpm','rad_s','rps'}
            assert not fault['labels_verified'] and not fault['speed_unit_verified']
            assert 'shift_tcn_L64_s17' not in models
            for model in fault['selected_models']:
                matched=df[df.run_id==model['run_id']].iloc[0]
                assert matched.artifact_sha256==model['artifact_sha256']
            fault_status='PASS 200 separate records, 13 frozen predictors + raw; unlabeled descriptive analysis'
        else:
            raise AssertionError('Fault study exists but is incomplete')
    nb=nbformat.read(root/'notebooks/00_COMPLETE_PROJECT_WALKTHROUGH.ipynb',as_version=4)
    nbformat.validate(nb)
    cells=[c for c in nb.cells if c.cell_type=='code']
    assert cells and all(c.execution_count is not None for c in cells)
    assert not any(o.output_type=='error' for c in cells for o in c.outputs)
    execution=json.loads((root/'notebooks/execution_summary.json').read_text())
    assert execution['sha256']==sha256(root/'notebooks/00_COMPLETE_PROJECT_WALKTHROUGH.ipynb')
    figures=pd.read_csv(root/'reports/figure_manifest.csv')
    for figure in figures.to_dict('records'):
        assert sha256(root/figure['path'])==figure['sha256']
    decks={}
    for name in ('Progress_Presentation_8min','Final_Presentation_12min'):
        prs=Presentation(root/'presentations'/(name+'.pptx'))
        assert 8<=len(prs.slides)<=20
        assert abs(prs.slide_width/prs.slide_height-16/9)<.01
        decks[name]=len(prs.slides)
    broken=[]
    anchor_issues=[]
    for folder in ('docs','reports','research'):
        for file in (root/folder).glob('*.md'):
            text=file.read_text()
            for link in re.findall(r'\]\(([^)]+)\)',text):
                if link.startswith(('http:','https:','mailto:')):continue
                target=link.split('#')[0].split(' "')[0]
                if target and not (file.parent/target).exists():
                    broken.append(dict(file=str(file.relative_to(root)),link=link))
                if '#' in link:
                    target_file=file.parent/target if target else file
                    fragment=link.split('#',1)[1]
                    if target_file.exists() and target_file.suffix=='.md' and fragment:
                        headings=[]
                        for heading in re.findall(r'^#{1,6}\s+(.+)$',target_file.read_text(),flags=re.M):
                            slug=re.sub(r'[^\w\- ]','',heading.lower()).replace(' ','-')
                            headings.append(slug)
                        if fragment not in headings:
                            anchor_issues.append(dict(file=str(file.relative_to(root)),link=link))
            if text.count('```')%2:
                raise AssertionError(f'Unbalanced Markdown code fences: {file}')
    if broken:
        (root/'reports/broken_links.json').write_text(json.dumps(broken,indent=2))
        raise AssertionError(f'{len(broken)} broken Markdown file links; see reports/broken_links.json')
    if anchor_issues:
        (root/'reports/broken_anchors.json').write_text(json.dumps(anchor_issues,indent=2))
        raise AssertionError(f'{len(anchor_issues)} broken Markdown anchors; see reports/broken_anchors.json')
    status=dict(status='PASS available-artifact integrity',runs=len(df),common_targets=int(df.n_targets.iloc[0]),
                executed_notebook_code_cells=len(cells),verified_figure_hashes=len(figures),decks=decks,broken_file_links=broken,
                real_fault_analysis=fault_status,anchor_full_text='BLOCKED network access',
                external_platform_execution='Windows/macOS unverified')
    (root/'reports/release_verification.json').write_text(json.dumps(status,indent=2))
    return status


def package(root, output):
    """Exclude raw user data, Git internals, caches and editable-install metadata."""
    root=Path(root); output=Path(output)
    def retained(path):
        rel=path.relative_to(root)
        return path.is_file() and not any(x in rel.parts for x in ('.git','__pycache__','.pytest_cache','raw','rendered')) and not any(x.endswith('.egg-info') for x in rel.parts) and path.suffix not in ('.zip','.pyc') and path.name!='checksums.sha256'
    files=sorted(p for p in root.rglob('*') if retained(p))
    hashes='\n'.join(sha256(p)+'  '+p.relative_to(root).as_posix() for p in files)+'\n'
    (root/'checksums.sha256').write_text(hashes)
    files.append(root/'checksums.sha256')
    with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for p in files:z.write(p,arcname='timeSeriesModel/'+p.relative_to(root).as_posix())
    with zipfile.ZipFile(output) as z:
        assert z.testzip() is None
        for p in files:
            content=z.read('timeSeriesModel/'+p.relative_to(root).as_posix())
            assert hashlib.sha256(content).hexdigest()==sha256(p)
    return dict(path=str(output),size_bytes=output.stat().st_size,sha256=sha256(output),files=len(files))
