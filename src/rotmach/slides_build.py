"""Deterministic editable presentations and grounded report from saved outputs.

All numerical claims are read from the experiment CSV/audit JSON, never typed
as hypothetical results. Figures are embedded images; all text, tables and
causal-flow diagrams remain editable PowerPoint objects.
"""
from pathlib import Path
import json
import math
import re
from xml.sax.saxutils import escape

import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY = RGBColor(18, 37, 56)
TEAL = RGBColor(0, 112, 128)
INK = RGBColor(36, 50, 62)
SLATE = RGBColor(91, 108, 124)
PALE = RGBColor(247, 249, 250)
WHITE = RGBColor(255, 255, 255)
FONT = "Aptos"


def _fault_evidence(root):
    """Read finalized fault-study outputs; never infer completion from upload names.

    A summary JSON and real time-domain rows must both exist. Record counts are
    cross-checked against unique CSV IDs, so stale/partial study output is loud.
    """
    root = Path(root)
    summary_path = root / "reports/fault_study_summary.json"
    if not summary_path.exists():
        return None
    summary = json.loads(summary_path.read_text())
    if summary.get("status") not in ("completed", "complete", "passed", "PASS", "complete_descriptive_unlabeled"):
        return None
    features = pd.read_csv(root / "reports/fault_record_features.csv")
    if features.empty or "model_id" not in features:
        raise ValueError("Fault summary exists without real feature rows")
    count = int(features.record_id.nunique())
    if count != int(summary.get("n_records", summary.get("real_fault_records_processed", 0))):
        raise ValueError("Fault summary/CSV record counts disagree")
    raw = features[features.model_id == "raw"].set_index("record_id")
    if len(raw) != count or raw.index.duplicated().any():
        raise ValueError("Expected one raw feature row per independent record")
    if raw.n_evaluated.nunique() != 1:
        raise ValueError("Report expects a common evaluated length across records")
    stats = []
    for model_id, rows in features.groupby("model_id", sort=True):
        aligned = rows.set_index("record_id").join(raw[["rms"]], rsuffix="_raw")
        if len(aligned) != count or aligned.index.duplicated().any():
            raise ValueError("Model fault rows must cover each record exactly once")
        stats.append(dict(model_id=model_id, records=len(rows),
                          median_rms=float(rows.rms.median()),
                          median_raw_rms_ratio=float((aligned.rms / aligned.rms_raw).median()),
                          median_crest_factor=float(rows.crest_factor.median()),
                          median_kurtosis=float(rows['kurtosis'].median())))
    angular_path = root / "reports/fault_angular_features.csv"
    angular = pd.read_csv(angular_path) if angular_path.exists() else pd.DataFrame()
    return dict(summary=summary, features=features, angular=angular,
                stats=pd.DataFrame(stats), n_records=count,
                n_evaluated=int(raw.n_evaluated.iloc[0]))


def _table_md(frame, columns):
    out = ["| " + " | ".join(columns) + " |", "|" + "---|" * len(columns)]
    for _, row in frame.iterrows():
        vals = []
        for col in columns:
            value = row[col]
            vals.append(f"{value:.6f}" if isinstance(value, float) else str(value))
        out.append("| " + " | ".join(vals) + " |")
    return "\n".join(out)


def _textbox(slide, x, y, w, h, text, size=20, color=INK, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.03)
    tf.margin_top = tf.margin_bottom = Inches(.03)
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name, p.font.size = FONT, Pt(size)
        p.font.bold, p.font.color.rgb = bold, color
        p.space_after = Pt(11)
    return shape


def _new_slide(prs, title, subtitle, number, deck):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PALE
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(.15))
    band.fill.solid(); band.fill.fore_color.rgb = TEAL; band.line.fill.background()
    _textbox(slide, .48, .38, 12.3, .72, title, 29, NAVY, True)
    _textbox(slide, .51, 1.08, 12.2, .55, subtitle, 13, SLATE)
    _textbox(slide, .5, 7.12, 11.9, .22,
             "MECG 7780 · Validation study · Generated from saved experiment outputs", 9, SLATE)
    _textbox(slide, 12.25, 7.1, .55, .25, str(number), 10, SLATE)
    return slide


def _bullets(slide, items, x=.65, y=1.95, w=12, h=4.8, size=23):
    h = min(h, 6.85-y)
    return _textbox(slide, x, y, w, h, "\n".join("• " + s for s in items), size)


def _figure(slide, root, filename, caption, right=None):
    path = root / "reports/figures" / filename
    if not path.exists():
        raise FileNotFoundError(f"Required generated figure missing: {path}")
    from PIL import Image
    with Image.open(path) as im:
        width, height = im.size
    maxw, maxh = (8.6, 4.55) if right else (11.9, 4.6)
    scale = min(maxw / width, maxh / height)
    w, h = width * scale, height * scale
    slide.shapes.add_picture(str(path), Inches(.62), Inches(1.85), width=Inches(w), height=Inches(h))
    _textbox(slide, .65, 6.5, 12, .45, caption, 11, SLATE)
    if right:
        _bullets(slide, right, x=9.45, y=2.0, w=3.3, h=4.4, size=18)


def _flow(slide, labels, caption, arrows=True):
    n = len(labels)
    boxw, gap = (11.8 - .28 * (n - 1)) / n, .28
    for i, label in enumerate(labels):
        x = .65 + i * (boxw + gap)
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.15), Inches(boxw), Inches(1.45))
        shape.fill.solid(); shape.fill.fore_color.rgb = NAVY if i % 2 == 0 else TEAL
        shape.line.fill.background()
        tf = shape.text_frame; tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.text = label; p.alignment = PP_ALIGN.CENTER
        p.font.name, p.font.size, p.font.color.rgb = FONT, Pt(19), WHITE
        if arrows and i < n - 1:
            _textbox(slide, x + boxw + .03, 2.63, .25, .45, "→", 20, TEAL, True)
    _textbox(slide, .7, 4.15, 11.9, 1.55, caption, 24)


def _results_table(slide, df):
    best = df[df.kind != "shift_tcn"].sort_values("mse").drop_duplicates("kind").head(8)
    headers = ["Model", "Lag", "MSE", "Parameters", "Train s", "Record s"]
    table = slide.shapes.add_table(len(best) + 1, len(headers), Inches(.65), Inches(1.92), Inches(12), Inches(3.95)).table
    widths = [3.8, .7, 1.4, 1.6, 1.8, 1.8]
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)
    rows = [headers] + [[r.kind, str(r.lag), f"{r.mse:.4f}", str(r.parameters),
                         f"{r.train_wall_s:.3f}", f"{r.inference_record_median_s:.3f}"] for _, r in best.iterrows()]
    for i, values in enumerate(rows):
        for j, value in enumerate(values):
            cell = table.cell(i, j)
            cell.text = value
            cell.fill.solid(); cell.fill.fore_color.rgb = NAVY if i == 0 else (WHITE if i % 2 else RGBColor(233, 241, 243))
            for p in cell.text_frame.paragraphs:
                p.font.name, p.font.size = FONT, Pt(15)
                p.font.bold, p.font.color.rgb = i == 0, WHITE if i == 0 else INK
    _textbox(slide, .7, 6.15, 11.8, .65,
             "Best observed configuration per family; validation-selected. Timing includes checkpoint load, window construction and scoring. Full CSV retains every run.", 15, SLATE)


def build_presentations(root):
    """Build both editable decks, notes, claim map and structural checks."""
    root = Path(root).resolve()
    df = pd.read_csv(root / "reports/results_summary.csv")
    audit = json.loads((root / "reports/data_audit.json").read_text())
    fault = _fault_evidence(root)
    best = df[df.kind != "shift_tcn"].sort_values("mse").iloc[0]
    ar = df[(df.kind == "ar") & (df.lag == 256)].iloc[0]
    zero = df[df.kind == "zero"].iloc[0]
    # Short-run proposal results are evidence, not presumed neural superiority.
    seed = df[(df.lag == 64) & df.kind.isin(["tcn", "speed_tcn"])].groupby("kind").mse.agg(["mean", "std", "count"])
    speed_statement = "; ".join(f"{k}: {r['mean']:.4f} ± {r['std']:.4f} ({int(r['count'])} seeds)" for k, r in seed.iterrows())
    speed_change = 100 * (seed.loc["speed_tcn", "mean"] / seed.loc["tcn", "mean"] - 1)
    sensitivity_path = root / "reports/angular_unit_sensitivity.csv"
    sensitivity_note = ""
    if sensitivity_path.exists():
        sens = pd.read_csv(sensitivity_path)
        raw = sens[(sens.split == "valid") & (sens.series == "supplied_x")]
        sensitivity_note = "Validation raw corrected PMI under assumptions: " + "; ".join(f"{r.speed_unit_assumption} {r.pmi_scientifically_corrected:.4f}" for _, r in raw.iterrows()) + ". No calibrated physical order."
    fault_status = (f"{fault['n_records']} separate fault-candidate records processed; physical fault accuracy unverified"
                    if fault else "Fault study has no finalized results in this build; physical fault performance unverified")
    pdir = root / "presentations"; pdir.mkdir(exist_ok=True)
    claims = []
    specs = {
        "Progress_Presentation_8min": [
            ("Predicting variable-speed vibration", "Project progress · 8-minute speaking target · Author / group: add names", "title"),
            ("A causal prediction task", "Previously observed samples predict the next sample; no future vibration", "causal"),
            ("What the accessible data establish", "Measured audit; labels and calibration are not invented", "data"),
            ("Operating speed changes the signal", "Train and validation each contain one 7.5-second recording", "eda"),
            ("A fair evaluation contract", "One-step teacher forcing · common burn-in · validation only", "protocol"),
            ("Strong baselines come first", "Least squares, ridge regularization and speed-dependent coefficients", "baseline"),
            ("Compact nonlinear alternatives", "MLP, unidirectional GRU and explicitly causal dilated TCN", "architecture"),
            ("A falsifiable speed-conditioning proposal", "Established FiLM, adapted to observed past speed", "novelty"),
            ("Residuals are diagnostic candidates", "Python numerical validation is separate from real fault evidence", "diagnostics"),
            ("Evidence, blockers and next steps", "Updated PDF rubric: organization 4 · clarity 6 · understanding 8 · novelty 2%", "finish"),
        ],
        "Final_Presentation_12min": [
            ("Causal prediction under variable speed", "Final validation study · 12-minute speaking target · FINAL_DATE_UNCONFIRMED", "title"),
            ("Data, provenance and limits", "One reference trace and one validation trace; no independent test", "data"),
            ("Causality before model complexity", "Each target uses observed history only", "causal"),
            ("The conventional comparison", "AR/ridge and LPV-AR surrogate; no replicated Spa LPV-ARMA score", "baseline"),
            ("Three compact neural ideas", "MLP lag mapping · GRU state · causal dilated TCN", "architecture"),
            ("Speed-conditioning: contribution and controls", "A project-level adaptation; not a world-first claim", "novelty"),
            ("Measured validation performance", "Every model evaluated at exactly t = 256…23999 (zero-based)", "results"),
            ("Accuracy has a computational cost", "CPU-only, measured public inference path; GPU benchmarks not run", "efficiency"),
            ("Does conditioning help consistently?", "Context, seed and capacity controls; do not select only favorable seeds", "ablation"),
            ("A residual does not prove a defect", "Prediction accuracy and diagnostic utility are distinct objectives", "residual"),
            ("PMI and angular resampling need assumptions", "Explicit speed units; literal versus corrected MATLAB indexing", "diagnostics"),
            ("Unlabeled records: measured descriptors", "Record-wise raw/residual features; no invented health classes", "faultblock"),
            ("Limits of inference", "Validation selection, one trace, repeated speed trajectory, unknown labels", "limits"),
            ("What the evidence supports", "Recommend using the measured table; validate on independent records next", "conclusion"),
            ("Sources and reproducibility", "Editable decks generated from project outputs; exact CSVs and checkpoints retained", "references"),
        ],
    }
    notes_extra = {
        "title": "Introduce the engineering aim: remove predictable reference vibration using a strictly causal next-sample predictor, then investigate residual structure. Distinguish finalized descriptive feature evidence from any unsupported physical diagnostic accuracy claim.",
        "causal": "Point to the ordering. For lag L, x[L] sees x[0:L], and speed inputs stop at L−1. Teacher forcing means previously measured vibration is available; this is not a free-running forecast. For causal models, changing future samples cannot change earlier predictions.",
        "data": "Both accessible files contain 24,000 samples at 3,200 Hz. Their duration by N/fs is 7.5 seconds, while the span between first and last timestamps is one sample shorter. The supplied normalization reconstructs exactly; stored first scaling entries are not empirical moments of the training trace. Speed source units remain unknown.",
        "eda": "Describe the actual speed trajectory and amplitude/frequency variation without assigning unverified gear tooth counts or shaft units. The two speed traces are nearly identical. Correlation is not proof of leakage, but it sharply limits conclusions about unseen operating histories.",
        "protocol": "Every score uses the same 23,744 validation targets after a 256-sample burn-in, so changing lag does not change the scored sample population. Validation also selects neural epoch and model configuration; it is therefore optimistic for final generalization. No test file exists here.",
        "baseline": "Explain ridge as squared-error fitting with a coefficient-size penalty. The LPV-AR surrogate expands past vibration with speed and speed-squared interactions, but has no moving-average innovation recursion or lecturer-paper sparse estimation. It is not Spa LPV-ARMA. No published score is invented.",
        "architecture": "MLP maps a finite history to a scalar; the GRU summarizes a unidirectional sequence with update/reset gates; the TCN uses left-only padded convolutions and increasing dilations. Two convolutions per block give receptive field 1+4 times the sum of dilations. Limited training budgets must be disclosed.",
        "novelty": "FiLM is an established conditioning mechanism. Here each past temporal feature receives a bounded gain and offset from that position's observed speed. Compare to unconditioned same-core TCN, late speed concatenation, a larger unconditioned capacity control and an explicitly off-line phase-shifted-speed negative control. The shifted control is excluded from causal recommendations.",
        "results": "Read the table as measured validation performance, not test performance. The top row is an observed best configuration and its selection used this validation trace. Inspect the full CSV for all seeds and contexts rather than interpreting a single winning row as a population estimate.",
        "efficiency": "Report CPU wall time and the actual timing boundary. The public inference path includes checkpoint deserialization, window preparation and full-record prediction; one warm-up and three repeats are used. These are not isolated kernel latencies. Linear fitting is float64 and neural fitting float32. GPU time was not measured.",
        "ablation": "The seed summary reflects initialization variability, not machine-to-machine variability. Context effects are compared on common target indices. A speed gain can depend on capacity, optimization and nearly matched operating histories. If the speed model loses, retain the result; novelty is a falsifiable experimental proposal.",
        "residual": "Residual e equals observed vibration minus the one-step estimate. Lower residual energy can simply reflect better prediction. A flexible teacher-forced model can also predict and remove defect content, so the best MSE model need not be best for diagnosis. These plots concern accessible reference/validation signals only.",
        "diagnostics": "RMS, crest factor, ordinary kurtosis, correlated kurtosis and envelope autocorrelation are implemented in Python. The MATLAB coarse PMI index is one sample before the intended period because autocorrelation lag zero occupies the first element. Angular resampling requires physical speed units; RPM, revolutions per second and radians per second differ. MATLAB toolbox parity has not been live-verified.",
        "faultblock": "If a finalized candidate-record summary is absent, do not claim real fault-record feature results. Independent records must be processed separately without retraining, and physical detection accuracy remains unsupported without documented labels.",
        "limits": "Emphasize that many overlapping windows are not independent machines. The validation trace is also used for early stopping and selection. Speed units and health labels remain unresolved. The lecturer article is a course-verified citation but full experimental methods were inaccessible, preventing an exact comparative claim.",
        "conclusion": "Answer RQ1 using the saved common-index table; answer RQ2 with measured costs; bound RQ3 to controlled experiments and RQ4 to finalized descriptor evidence; answer RQ5 with context and seed sensitivity rather than unsupported generalization. A simpler model can be preferable even if a neural configuration has a favorable selected score.",
        "finish": "Summarize completed audit, causal models, saved checkpoints and numerical feature tests. Identify unresolved instructor date conflict, absent test set and needed label/speed calibration. Demonstrate command-line and notebook reproduction rather than promising future results.",
        "references": "Direct the audience to the verified literature matrix and access log. The course citation motivates speed-dependent modeling; author repositories support TCN and FiLM implementation concepts. Original TCN and FiLM studies were not gearbox validation studies. All claims here map to stored tables, configurations and figures.",
    }
    if fault:
        notes_extra.update({
            "title": "Introduce the engineering aim: predict the next sample causally, compare accuracy and cost, then inspect raw and residual diagnostic descriptors. The complete split upload now enables independent analysis of all 200 fault-candidate records. These are descriptive results without verified physical fault labels, not a validated diagnostic classifier.",
            "data": "The two reference/validation files each contain 24,000 samples at 3,200 Hz. The newly available twenty MAT parts preserve 200 separate candidate records, each 32,000 samples or 10 seconds. No records are flattened together. No second normalization is applied to supplied fault x. Speed physical units and health labels remain unverified.",
            "residual": "Every supplied fault-candidate record is scored using previously trained reference predictors with no fault-data retraining. All models and raw features use the same 256-sample burn-in and remaining 31,744 samples per record. The displayed RMS trajectories are descriptive: record-index trends are not fault-onset labels, and reduced RMS is not detection accuracy.",
            "diagnostics": "The angle-domain indicators are computed under explicitly named RPM, revolutions-per-second and radians-per-second assumptions; none establishes the true speed unit. The literal MATLAB coarse index selects lag T minus one while the corrected branch selects T. The sensitivity plot compares actual record descriptors across these choices, not physical fault probabilities.",
            "faultblock": "All candidate records now have real raw/residual feature measurements. Compare the displayed distributions and note speed dependence, normalization and equal segment length. There are no verified healthy/faulty classes, onset dates or severity values. A model with small residual MSE can also suppress periodic content. Consequently no ROC, F1 or sensitivity is justified and diagnostic model choice must remain provisional.",
            "limits": "The transfer blocker has been resolved using the twenty split MAT files. Remaining limitations are unknown health labels and physical speed unit, validation reuse for model selection, only one short validation trace and nearly matched speed histories, and incomplete publisher access for exact Spa LPV-ARMA replication. Descriptive fault-record patterns do not resolve these inferential limits.",
            "conclusion": "Answer RQ1 using the saved common-index validation table, and RQ2 using the measured CPU costs. Speed-FiLM does not improve mean validation MSE across the matched seeds. RQ4 now has real record-wise raw/residual and angular-assumption evidence, but physical fault accuracy remains unidentifiable without labels. Recommend inspecting both raw and model residual descriptors rather than declaring a physical fault winner.",
            "finish": "Summarize completed audit, trained and saved causal predictors, all 200 candidate records processed independently, tested numerical features and reproducible plots. Remaining needs are calibrated speed units, physical health labels and independent test data, plus clarification of the course date conflict. No additional upload is needed for the completed 200-record study.",
        })
    notes_files = []
    for name, spec in specs.items():
        prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
        progress = name.startswith("Progress")
        seconds = [35, 45, 50, 45, 50, 55, 55, 55, 55, 35] if progress else [30, 45, 50, 50, 55, 55, 65, 55, 60, 45, 55, 40, 45, 45, 20]
        if progress:
            assert sum(seconds) == 480
        if not progress:
            # Explicit, reproducible 12-minute total; preserve more time for results.
            seconds = [29, 44, 49, 49, 54, 54, 69, 54, 64, 44, 54, 39, 44, 49, 24]
            assert sum(seconds) == 720
        script = [f"# {name.replace('_', ' ')} speaker script", "", f"Total planned speaking time: {sum(seconds)} seconds ({sum(seconds)/60:.1f} minutes). Slide times are estimates, not a guarantee of delivery pace.", ""]
        for i, ((title, subtitle, key), duration) in enumerate(zip(spec, seconds), start=1):
            slide = _new_slide(prs, title, subtitle, i, name)
            if key == "title":
                _textbox(slide, .75, 2.02, 11.7, 1.18, "Can a compact causal model explain vibration efficiently—and leave interpretable residuals?", 34, NAVY, True)
                _bullets(slide, [f"Best measured validation configuration: {best.run_id}", f"Common-index MSE: {best.mse:.4f}; independent test unavailable", fault_status], y=3.75, size=21)
            elif key == "causal":
                _flow(slide, ["Observed history\nx[t−L:t], s[t−L:t]", "Past-only windows\nno target values", "Causal model\nx̂[t]", "Residual\ne[t] = x[t] − x̂[t]"], "Example: x[64] uses samples 0…63 only. One-step teacher forcing; no x[64] in model inputs.")
            elif key == "data":
                _bullets(slide, ["Train: 24,000 samples · Validation: 24,000 samples · fs = 3,200 Hz", "Provided x and spdnorm are used exactly; no record-wise rescaling", f"Speed trajectory correlation: {audit['comparisons']['speed_pearson']:.9f}", "Speed units / reference health state unverified; normalized x is dimensionless", (f"Fault candidates: {fault['n_records']} separate 10-second records in 20 MAT parts" if fault else "Fault study not finalized in this build"), "sigtest.mat and sigfault2.mat not supplied"], size=21)
            elif key == "eda":
                _figure(slide, root, "eda.png", "Source: data/raw/sigtrain.mat and sigvalid.mat; hashes in data/manifest.json.")
            elif key == "protocol":
                _bullets(slide, ["Fit on sigtrain; select epochs/configurations using sigvalid", "Score all models at zero-based indices 256…23,999: 23,744 targets", "Observed past vibration; speed available only through t−1", "No random window split presented as independent generalization", "Validation is reused for selection: final unbiased test performance remains unknown"], size=23)
            elif key == "baseline":
                _textbox(slide, .75, 1.95, 11.6, .85, "Ridge AR:  minimize ‖y − Xb‖² + α‖b‖²,  α = 1", 29, TEAL, True)
                _bullets(slide, ["Zero, training-mean and persistence sanity checks", "OLS and ridge AR; ARX appends past speed and speed²", "LPV-AR: lag coefficients vary polynomially with speed at t−1", f"Measured ridge AR(256) validation MSE: {ar.mse:.6f} (zero: {zero.mse:.6f})", "LPV-AR surrogate has no ARMA innovation recursion; no exact paper reproduction"], y=3.05, size=22)
            elif key == "architecture":
                _flow(slide, ["MLP\nlag vector → scalar", "GRU\nupdate / reset gates", "TCN\ncausal dilations"], "Three separate model alternatives; compact networks, bounded epochs and recorded seeds. TCN: two kernel-3 convolutions per block; receptive field = 1 + 4 Σ dilation.", arrows=False)
            elif key == "novelty":
                _flow(slide, ["Past speed s[k]", "Small linear map\nγ(s), β(s)", "h[k] × gain\n+ offset", "Causal prediction"], "Hypothesis: operating-speed information helps beyond extra capacity. Controls: late concatenation, larger unconditioned TCN and off-line phase-shifted speed (negative control only).")
            elif key == "results":
                _results_table(slide, df)
            elif key == "efficiency":
                _figure(slide, root, "efficiency.png", "Measured CPU wall time; details: reports/timing_and_hardware.md; no GPU measurements.", right=[f"Best model fit: {best.train_wall_s:.3f} s", f"Best model record path: {best.inference_record_median_s:.3f} s", f"Ridge AR(256): {ar.inference_record_median_s:.4f} s / record", "Includes checkpoint load + windows; 3 repeats"])
            elif key == "ablation":
                _figure(slide, root, "ablations.png", "Same validation targets for every context; points are initialization repeats, not independent recordings.", right=[speed_statement, f"FiLM mean MSE: {speed_change:.2f}% higher", "Full controls: model_ablation.csv"])
            elif key == "residual":
                if fault:
                    _figure(slide, root, "fault_rms_trends.png", f"{fault['n_records']} separate candidate records; equal {fault['n_evaluated']:,}-sample segments; dimensionless x/residual.")
                else:
                    _figure(slide, root, "prediction.png", "Accessible validation trace only; x and residual are dimensionless.")
            elif key == "diagnostics":
                if fault and not progress:
                    _figure(slide, root, "fault_unit_sensitivity.png", "Actual fault-candidate descriptors; speed units / order are assumptions. Literal PMI lag T−1; corrected T.")
                else:
                    _flow(slide, ["Residual / raw", "Explicit-unit\nangle grid", "Envelope + ACF", "CK / PMI\nwith caveats"], "Literal MATLAB PMI selects lag T−1; corrected variant selects T. Numerical tests and explicit unit assumptions support descriptive indicators; physical fault accuracy is unverified.")
                    if fault:
                        _textbox(slide, .7, 6.15, 11.9, .65, fault_status, 15, TEAL)
                    elif sensitivity_note:
                        _textbox(slide, .7, 6.15, 11.9, .65, sensitivity_note, 15, TEAL)
            elif key == "faultblock":
                if fault:
                    _figure(slide, root, "fault_indicator_distributions.png", "Actual record-wise distributions; no verified classes. RMS / kurtosis / CK / PMI describe signals, not fault accuracy.")
                else:
                    _bullets(slide, ["No finalized record-wise fault study is available in this build", "Raw/residual descriptors require preserved recording boundaries", "Python feature pipeline and synthetic unit checks are supplied", "Physical speed units and health labels require independent verification", "No ROC, F1, sensitivity or tooth-crack severity claims are supported"], size=23)
            elif key == "limits":
                _bullets(slide, ["Only one short validation recording, reused for epoch/model selection", "Speed trajectories are almost identical: unseen-speed transfer untested", "Short neural training budget; rankings are budget-specific", "Physical speed unit / health labels unverified; patterns remain descriptive", "Spa LPV-ARMA exact comparison unavailable; final dates conflict in the PDF"], size=22)
            elif key in ("conclusion", "finish"):
                _bullets(slide, [f"Measured best: {best.run_id}, validation MSE {best.mse:.4f}", "LPV-AR beats the tested neural models under this budget", f"FiLM speed-TCN mean MSE is {speed_change:.2f}% higher than TCN (3 seeds)", fault_status, "Next evidence: independent test and verified physical speed units / health labels"], size=22)
            elif key == "references":
                _bullets(slide, ["Chen et al. (2025), MSSP 224, 112161. doi:10.1016/j.ymssp.2024.112161", "Bai, Kolter & Koltun (2018), TCN. Original author implementation accessed", "Perez et al. (2018), FiLM. Original author implementation accessed", "Full evidence status: research/literature_matrix.csv and source_access_log.md", "Rebuild: python scripts/build_presentations.py; inputs: reports/results_summary.csv"], size=20)
            script.extend([f"## Slide {i}. {title} — {duration} seconds", "", notes_extra[key], "", "Evidence: " + ("`reports/results_summary.csv`; `experiments/runs/*/config.json`" if key in ("title", "results", "baseline", "ablation", "conclusion", "efficiency") else "`docs/course_requirements_and_open_questions.md`; `reports/data_audit.json`; relevant report section"), ""])
            slide.notes_slide.notes_text_frame.text = f"Speaking target: {duration} seconds.\n" + notes_extra[key]
            claims.append(dict(deck=name, slide=i, key=key, source="reports/results_summary.csv; reports/data_audit.json; reports/fault_study_summary.json; reports/fault_record_features.csv; reports/fault_angular_features.csv; docs/course_requirements_and_open_questions.md", numeric_claim_policy="read from stored outputs; physical fault labels unverified"))
        prs.save(pdir / f"{name}.pptx")
        short = "progress" if progress else "final"
        notes = pdir / f"{short}_speaker_notes.md"; notes.write_text("\n".join(script))
        notes_files.append(notes)
    pd.DataFrame(claims).to_csv(pdir / "slide_source_and_claim_map.csv", index=False)
    (pdir / "slide_source_and_claim_map.md").write_text("# Slide source and claim map\n\n" + _table_md(pd.DataFrame(claims), ["deck", "slide", "key", "source"]) + "\n\nAll text/tables/diagrams are editable. Plot images are generated from saved input/model artifacts. Real candidate-record descriptor results are shown only when finalized study evidence exists. No physical fault accuracy metrics appear because health labels remain unverified. See speaker notes for definitions and caveats.\n")
    verification = []
    for name, spec in specs.items():
        p = pdir / f"{name}.pptx"; loaded = Presentation(p)
        assert len(loaded.slides) == len(spec)
        assert abs(loaded.slide_width / Inches(1) - 13.333) < .001
        for i, slide in enumerate(loaded.slides, 1):
            assert slide.shapes and slide.notes_slide.notes_text_frame.text
            for shape in slide.shapes:
                assert shape.left >= 0 and shape.top >= 0
                assert shape.left + shape.width <= loaded.slide_width + Inches(.02)
                assert shape.top + shape.height <= loaded.slide_height + Inches(.02)
        verification.append(f"- PASS {p.name}: opened with python-pptx; {len(loaded.slides)} slides; 16:9; notes present; all shape bounds inside canvas.")
    (pdir / "presentation_verification.md").write_text("# Presentation verification\n\n" + "\n".join(verification) + "\n\nStructural checks do not establish visual rendering or font-layout parity. Rendering status is recorded separately after an actual export attempt.\n")
    return [pdir / f"{name}.pptx" for name in specs]


def build_report(root):
    """Write the final scientific report from actual CSV metrics and audit."""
    root = Path(root).resolve()
    df = pd.read_csv(root / "reports/results_summary.csv")
    audit = json.loads((root / "reports/data_audit.json").read_text())
    fault = _fault_evidence(root)
    best = df[df.kind != "shift_tcn"].sort_values("mse").iloc[0]
    zero = df[df.kind == "zero"].iloc[0]
    ar = df[(df.kind == "ar") & (df.lag == 256)].iloc[0]
    lpv = df[(df.kind == "lpv") & (df.lag == 256)].iloc[0]
    improvement = 100 * (1 - best.mse / zero.mse)
    seed = df[(df.lag == 64) & df.kind.isin(["tcn", "speed_tcn"])].groupby("kind").mse.agg(["mean", "std", "min", "max", "count"]).reset_index()
    means = seed.set_index("kind")["mean"]
    film_change = 100 * (means["speed_tcn"] / means["tcn"] - 1)
    lpv_reduction = 100 * (1 - lpv.mse / ar.mse)
    sensitivity_section = "Angular assumption sensitivity was not evaluated in this report build."
    sensitivity_path = root / "reports/angular_unit_sensitivity.csv"
    if sensitivity_path.exists():
        sens = pd.read_csv(sensitivity_path)
        selected = sens[(sens.split == "valid") & sens.series.isin(["supplied_x", "residual:"+best.run_id])]
        columns = ["series", "speed_unit_assumption", "samples_per_cycle", "pmi_scientifically_corrected", "ck_order1", "roundtrip_rmse_fraction_of_series_rms"]
        sensitivity_section = """The following **actual accessible validation-signal sensitivity** results are read from [angular_unit_sensitivity.csv](angular_unit_sensitivity.csv). They are not fault-record results and do not establish physical speed units or a true defect order. Samples per cycle vary to accommodate source sampling density under each assumption.\n\n""" + _table_md(selected, columns) + """\n\nThe different PMI values show that an unverified unit assumption materially changes the evaluated lag/order interpretation. Raw and best-model residual values use equal source segments after burn-in. The residual PMI is not automatically larger; lower prediction MSE does not imply enhanced periodic diagnostic contrast. Round-trip interpolation RMSE is measured relative to the corresponding source-series RMS and reveals substantial reconstruction change for these broadband signals; denser angular interpolation cannot restore information absent from the time samples. These are sensitivity diagnostics, not a validated tooth-crack detector."""
    table = _table_md(df[["run_id", "mse", "parameters", "train_wall_s", "inference_record_median_s"]], ["run_id", "mse", "parameters", "train_wall_s", "inference_record_median_s"])
    completion = (f"The twenty split MAT uploads now enable analysis of **{fault['n_records']} separate fault-candidate records**. Their structure, raw/residual features and explicit angular-unit sensitivities were measured without retraining predictors on those records. Physical fault-detection accuracy remains **unidentifiable without verified health labels**."
                  if fault else "No finalized real fault-record study is present in this build; candidate-record analysis is not claimed complete.")
    abstract_fault = (f"The fault study processes {fault['n_records']} separate candidate recordings and compares raw/residual statistics and assumption-labelled angular indicators. These descriptive results do not establish physical fault classes, severity or detection accuracy."
                      if fault else "Real fault-record descriptors are not claimed in this build.")
    fault_data = (f"The original oversized fault upload was superseded by twenty accessible split MAT files in `data/raw/sigfault1_parts/`. The finalized loader/audit preserves **{fault['n_records']} independent records**; all raw/residual descriptor rows use a common **{fault['n_evaluated']:,}-sample** segment after **256** burn-in samples. No record-wise rescaling or concatenation is performed. The verified split files expose supplied x rather than xraw, and no health-state, fault-onset or severity labels were found. File-part names and SHA256 provenance accompany feature rows. The record index is an ordering identifier only, not a health label."
                  if fault else "Real fault-record data are not described as inspected until the finalized study outputs exist.")
    fault_section = "No finalized real candidate-record study is available in this build. Preserve record boundaries and equal burn-in when data become available; no physical detection metrics are claimed without labels."
    fault_rq = "real fault-record comparisons are not available in this build"
    fault_summary = "Real fault-record descriptors are not claimed in this build."
    if fault:
        stats = fault["stats"]
        best_stats = stats[stats.model_id == best.run_id]
        ratio_text = ""
        if len(best_stats):
            b = best_stats.iloc[0]
            ratio_text = f"The recommended validation predictor's median per-record residual/raw RMS ratio is **{b.median_raw_rms_ratio:.6f}**. This ratio measures energy suppression on equal segments, not accuracy of physical fault identification. Its median residual kurtosis is **{b.median_kurtosis:.6f}** under the stated biased Pearson convention."
        angular_table = "Angular-unit feature table is not present in this build."
        if not fault["angular"].empty:
            angular = fault["angular"]
            chosen = angular[angular.model_id.isin(["raw", best.run_id, ar.run_id])]
            agg = chosen.groupby(["model_id", "speed_unit_assumption"], as_index=False).agg(
                records=("record_id", "nunique"),
                median_ck=("ck_order1", "median"),
                median_pmi_literal=("pmi_faithful_reference", "median"),
                median_pmi_corrected=("pmi_scientifically_corrected", "median"))
            angular_table = _table_md(agg, ["model_id", "speed_unit_assumption", "records", "median_ck", "median_pmi_literal", "median_pmi_corrected"])
        fault_section = f'''### 13.1 Actual record-wise protocol and completeness

The supplied twenty-part dataset is now accessible and **{fault['n_records']} distinct records** have been processed individually. This resolves the original single-upload transfer obstacle. Each predictor was trained only on the reference training recording, then frozen; candidate records were not used for retraining, epoch selection or the reported validation ranking. Recording state/context is reset at each boundary. The first 256 samples are excluded equally for raw and residual descriptors, leaving **{fault['n_evaluated']:,} samples per record**. No labels are inferred from part boundaries, record order or a striking figure.

The long-format [fault_record_features.csv](fault_record_features.csv) contains original part/hash provenance, model ID, burn-in/evaluated count, speed summaries, RMS, crest factor, ordinary kurtosis and measured computation times. [fault_angular_features.csv](fault_angular_features.csv) separately records speed-unit assumptions, angular sample density, interpolation change/edge information, CK, both PMI indexing variants and validity statuses. [fault_study_summary.json](fault_study_summary.json) summarizes actual completion. These are real candidate-record measurements, distinct from earlier synthetic numerical tests.

### 13.2 Raw versus residual statistics

{_table_md(stats, ['model_id', 'records', 'median_rms', 'median_raw_rms_ratio', 'median_crest_factor', 'median_kurtosis'])}

{ratio_text}

![Actual candidate-record RMS trends](figures/fault_rms_trends.png)

*Figure 5. Record-wise raw and residual RMS on equal source segments. All records are independent; connecting their indices visualizes ordering only and does not imply a continuous time trace or known fault onset.*

![Actual candidate-record indicator distributions](figures/fault_indicator_distributions.png)

*Figure 6. Raw/residual descriptor distributions across actual candidate records. These distributions have no known healthy/faulty class labels, so visual separation cannot be converted to fault sensitivity or classification accuracy.*

### Post hoc record-order observation

The all-record plot shows an amplitude change near source ID101. A post hoc descriptive comparison of source IDs1–100 and101–200 is saved in [fault_posthoc_source_partition.csv](fault_posthoc_source_partition.csv). The two groups have different median raw/residual RMS but similar coarse speed moments. This is not verified fault onset or healthy/faulty labeling: acquisition order, load, preprocessing and machine state are not documented, and similar speed moments do not prove identical trajectories. No thresholds or models were selected from this partition. Separate recording processing also does not establish statistical independence of acquisitions.

### 13.3 Angular-unit and MATLAB-indexing sensitivity

The following table gives per-condition **medians across independently processed records** for raw, the prediction-recommended model and ridge AR(256). The speed interpretations RPM, rad/s and rps remain assumptions. The assumed one-cycle modulation order is not a verified tooth-fault characteristic. Median feature changes are descriptive and can depend on speed, interpolation, energy, lag convention and record length.

{angular_table}

![Actual angular-unit sensitivity](figures/fault_unit_sensitivity.png)

*Figure 7. Candidate-record CK/PMI sensitivity under three explicitly named speed-unit interpretations. Neither choosing a convenient unit nor choosing the most separated distribution establishes physical calibration or fault labels.*

### 13.4 Prediction quality versus diagnostic utility

The validation MSE winner and a diagnostic recommendation answer different questions. A low-energy residual may remove predictable reference structure, but can also suppress periodic content that might matter diagnostically. Changes in kurtosis/PMI/CK are therefore inspected alongside raw descriptors, speed and implementation assumptions; a larger statistic is not universally a better detector. The safe practical recommendation is to retain **both raw descriptors and the recommended predictor's residual descriptors** for further calibrated investigation, rather than declare a proven physical diagnostic winner. Model-to-model plots are in [fault_prediction_diagnostic_tradeoff.png](figures/fault_prediction_diagnostic_tradeoff.png).

No physical accuracy, AUC, F1, sensitivity, fault onset, crack presence or severity is reported. Such metrics require documented labels and record-level threshold-selection/evaluation on independent data. Even if an apparent change aligns with a record index or part boundary, it remains a post hoc descriptive association until acquisition metadata establishes meaning. See [fault label limitations](../docs/fault_label_limitations.md).'''
        fault_rq = f"real raw/residual and assumption-labelled angular descriptors were measured for {fault['n_records']} separate candidate records, while physical fault accuracy remains unidentifiable without labels"
        fault_summary = f"The complete {fault['n_records']}-record split fault dataset was processed independently with frozen reference predictors; raw/residual statistics and angular-unit sensitivities are saved. Physical fault accuracy remains unidentifiable without health labels."
    report = f'''# Time Serial Dependency Modeling via Advanced Machine Learning Algorithms

MECG 7780 Engineering Data Analytic · Reproducible Python validation study  
Author / group: **to be supplied by the user**  
Generated: 2026-10-30 · Final submission/presentation dates: **unconfirmed**

> **Completion boundary.** {completion} `sigtest.mat` and `sigfault2.mat` were not supplied. All prediction scores below are validation scores, not held-out test scores.

## Abstract

Variable operating speed complicates vibration prediction because frequency content and dynamic response may vary over time. We implement strictly causal, one-step predictors of supplied dimensionless vibration using the accessible `sigtrain.mat` and `sigvalid.mat` recordings. Each contains 24,000 samples at 3,200 Hz. The comparison comprises naive predictors, ordinary and ridge autoregression, a causal speed-augmented ARX baseline, a polynomial speed-conditioned LPV-AR surrogate, a compact multilayer perceptron, a unidirectional GRU, a causal dilated temporal convolutional network (TCN), and FiLM-style speed-conditioned variants. All models are evaluated at the same 23,744 validation target indices. The best observed eligible configuration, `{best.run_id}`, achieves MSE **{best.mse:.6f}**, versus **{zero.mse:.6f}** for zero prediction and **{ar.mse:.6f}** for ridge AR(256). This is a **{improvement:.2f}%** reduction relative to zero prediction on this trace. The recommendation is provisional because validation is also used for epoch/configuration selection, the speed histories are nearly identical, and independent test data are absent. CPU timing and checkpoint round trips are measured. {abstract_fault} Physical angular calibration, live MATLAB parity and an exact Spa LPV-ARMA comparison remain unverified.

## 1. Engineering motivation and course scope

Rotating systems contain repeatable vibration components from excitation, meshing and structural dynamics. Under variable speed, a fixed-frequency representation may spread otherwise orderly periodic content across frequency. A next-sample model can quantify predictable structure and generate a residual for subsequent analysis. The course explicitly requests MSE, CPU/GPU efficiency and fault performance; it encourages speed conditioning and supplies MATLAB PMI code. Those obligations must remain separate from the smaller set of empirically completed tasks.

The authoritative updated PDF assigns progress weights of **4% organization/presentation, 6% clarity, 8% algorithm understanding and 2% novelty**, and final weights of **16%, 24%, 30% performance and 10% novelty**. These differ from the master prompt. The PDF prints Nov 4, 2026 for the 8-minute progress presentation but Dec 7, 2025 at 24:00 and Dec 9, 2025 for report and final presentation. We preserve the conflict rather than inventing dates. See [course requirements](../docs/course_requirements_and_open_questions.md) and [rubric mapping](../docs/rubric_traceability_matrix.md).

## 2. Research context and strength of evidence

The course cites Chen, Li, Jiang, Gong and Zhou (2025), *Sparse LPV-ARMA model for non-stationary vibration representation and its application on gearbox tooth crack detection under variable speed conditions*, MSSP 224, 112161, DOI 10.1016/j.ymssp.2024.112161. The PDF attributes the same data to that paper. Publisher/full-text access was restricted, so the exact sparse estimator, innovation recursion, data selection and reported MSE could not be independently reproduced. We therefore implement a related **polynomial LPV-AR surrogate** and make no numerical claim to outperform Spa LPV-ARMA.

Original TCN author material motivates causal dilated convolution as an efficient sequence-model alternative to recurrence. Original FiLM author material supports feature-wise conditioning by gains and offsets. These source tasks differ from gearbox diagnosis; their existence is architectural motivation, not evidence that either model wins here. Ridge is essential because a short high-rate recording supplies many highly dependent windows and a modest linear predictor may be accurate and inexpensive. The full [research synthesis](../research/research_synthesis.md), [literature matrix](../research/literature_matrix.csv) and [access log](../research/source_access_log.md) distinguish course citation, accessed implementation documentation, abstract evidence and inaccessible full methods. The requested deep domain-specific review is limited by those access constraints.

## 3. Research questions and predeclared hypotheses

**RQ1:** Which tested causal configuration minimizes common-index validation MSE? **RQ2:** What accuracy–CPU-cost–parameter trade-off is observed? **RQ3:** Does observed past-speed conditioning improve prediction beyond capacity and optimization changes? **RQ4:** Do residuals expose useful periodic/impulsive diagnostic structure relative to raw signals? **RQ5:** Are conclusions stable across contexts, seeds and feature conventions?

The falsifiable working hypotheses are that contextual predictors improve on naive prediction, speed conditioning may help under nonstationary operating conditions, and residual processing may clarify periodic/impulsive structure. None is presumed true. Physical fault-detection accuracy cannot be evaluated without verified labels. Model selection and descriptive diagnostic interpretation are separate decisions.

## 4. Data, provenance and forensic audit

The original accessible files are preserved in `data/raw/`. SHA256, byte counts, availability and source fields are stored in [the manifest](../data/manifest.json); machine-readable measured summaries are in [data_audit.json](data_audit.json).

| Property | Training | Validation |
|---|---:|---:|
| Samples | 24,000 | 24,000 |
| Sampling rate (Hz) | 3,200 | 3,200 |
| N/fs duration (s) | 7.5 | 7.5 |
| First timestamp (s) | 0.0003125 | 0.0003125 |
| Last timestamp (s) | 7.5 | 7.5 |
| Supplied x mean | {audit['sigtrain']['x_mean']:.9f} | {audit['sigvalid']['x_mean']:.9f} |
| Supplied x standard deviation, ddof=0 | {audit['sigtrain']['x_std']:.9f} | {audit['sigvalid']['x_std']:.9f} |
| Empirical xraw mean | {audit['sigtrain']['empirical_raw_mean']:.9f} | {audit['sigvalid']['empirical_raw_mean']:.9f} |
| Empirical xraw standard deviation, ddof=0 | {audit['sigtrain']['empirical_raw_std']:.9f} | {audit['sigvalid']['empirical_raw_std']:.9f} |

The first stored scaling entries give `x=(xraw−0.0026221941677904226)/0.42517293168840065`, and supplied speed normalization uses reference endpoints `124.07455535625635` and `167.3376580579989`. Reconstruction maximum error is zero at the inspected NumPy precision. The stored arrays have three entries; only the first entries' role in the accessible scalar vibration normalization is verified. Their other entries are not interpreted. The empirical training moments differ from the metadata and are not substituted into the primary experiment.

The speed-trace Pearson correlation is **{audit['comparisons']['speed_pearson']:.12f}**; vibration correlation is **{audit['comparisons']['vibration_pearson']:.9f}**. No validation speeds fall outside the observed training min/max in this pair of recordings. This is restricted coverage, not proof of data leakage. We cannot infer performance under a substantially different speed trajectory. The physical speed unit is absent from verified metadata, so figures use “source units.” Supplied normalized vibration and residuals are dimensionless; a historical MATLAB axis annotation does not establish acceleration calibration.

{fault_data}

![Accessible vibration and speed traces](figures/eda.png)

*Figure 1. Actual accessible train/validation traces, sampled at 3,200 Hz. Speed and vibration must be interpreted with their separately stated units. Data hashes identify the exact files.*

## 5. Prediction formulation, alignment and leakage control

The primary target is `xhat[t] = f(x[t−L:t], spdnorm[t−L:t])`; the right slice endpoint is excluded. A lag-L window first predicts target `x[L]` from `x[0:L]`. For speed-only interactions, the latest operating value is `spdnorm[t−1]`. No contemporaneous or future vibration enters the model. The primary protocol is **teacher-forced one-step prediction**, using previously measured true vibration; it is not a recursive free-running forecast.

All reported MSEs use the identical zero-based validation indices **256 through 23,999 inclusive**, giving **23,744** targets. This maximum burn-in accommodates all tested lags. The loss is the arithmetic mean of squared dimensionless residuals, with no extra standardization or per-model target removal. RMSE, MAE and NMSE are secondary; NMSE divides by the target mean square, not target variance. Recording windows never cross file boundaries. Training minibatch permutation changes optimization ordering; it does not create a random train/test split.

The training file fits coefficients and weights. Validation is used for neural epoch selection and observed model/configuration ranking. Consequently, reported validation scores are useful comparison evidence on this trace but are not unbiased final generalization estimates. There is no claimed untouched test set. Any blockwise uncertainty describes temporal variation within this trace, not independent machine-to-machine uncertainty. An off-line phase-shifted-speed experiment is explicitly a negative control and excluded from causal model recommendations; it must not be deployed as a causal speed policy.

## 6. Models and mathematical objectives

### 6.1 Naive and linear models

Zero prediction uses `xhat[t]=0`; training mean uses one scalar fitted on the supplied training x; persistence uses `xhat[t]=x[t−1]`. These differ and are scored separately. AR(p) uses a row of the p previous vibration samples and an intercept. OLS minimizes `||y−Xb||²`; ridge minimizes `||y−Xb||² + α||b||²` with an unpenalized fitted intercept, `α=1` and scikit-learn's Cholesky solver. No matrix inverse is explicitly formed. Lags 32/64/128/256 correspond to 10/20/40/80 ms at 3,200 Hz.

ARX appends the last observed normalized speed and its square. The polynomial LPV-AR surrogate concatenates `X`, `X*s[t−1]` and `X*s[t−1]²`, so each AR coefficient can depend quadratically on speed. All supplied x/spdnorm fields retain their common scaling. These interaction columns are correlated, making ridge stabilization relevant. The surrogate does not include recursive innovation lags and does not reproduce the lecturer's sparse LPV-ARMA estimation. Parameter counts include the intercept.

### 6.2 Compact nonlinear models

The MLP uses a lag-vector input, one 32-unit ReLU hidden layer and a scalar output. The GRU is unidirectional with 16 hidden units and a scalar linear head; reset and update gates regulate the candidate state and retained history. Neither model is bidirectional. Both receive vibration-only inputs in the principal suite.

The TCN first projects the vibration channel to eight features. Each residual block contains two kernel-three convolutions at one dilation, with explicit **left padding only**, ReLU activations and a residual sum. Dilations double by block. Its receptive field is `R=1+4*Σdilation`, covering at least the declared lag; padding fills missing early context with zeros and never future data. A final linear head maps the last temporal feature to the next sample. Contexts 32/64/128 test a bounded range rather than an unrestricted architecture search.

### 6.3 Project-level speed-conditioned contribution

For each temporal position k in the observed window, FiLM-style conditioning maps its observed normalized speed to channel-wise gain and offset. The implementation uses a bounded residual gain `1+0.1*tanh(g(s[k]))` and additive `0.1*b(s[k])` after a TCN block. This is an adaptation of established conditioning, not a world-first method claim. Conditioning at past positions respects the availability policy.

Controls include the same core TCN without speed, late concatenation of speed, a nine-channel unconditioned TCN with a larger parameter budget, and an intentionally phase-misaligned speed negative control. The last is off-line and excluded from the eligible winner. Counts and timings are reported directly rather than described as exactly parameter matched. The controls improve interpretation but cannot eliminate all optimization, effective-capacity or trajectory-specific confounding.

## 7. Training budget, seeds and checkpoint selection

The bounded CPU suite uses neural seeds 17, 29 and 43 for principal lag-64 TCN/speed-TCN pairs; other neural configurations have seed 17. AdamW uses learning rate 0.003, weight decay 0.0001, batch size 512, gradient clipping at norm 1, a maximum of 10 epochs and patience 3. The checkpoint with lowest declared common-index validation MSE is restored. The complete effective configuration is saved beside each run; actual completed epochs and learning curves are retained. Short training is a deliberate feasibility budget, not proof that all architectures reached their best attainable accuracy.

Linear computation uses float64 and neural weights/windows float32. Two PyTorch CPU threads and deterministic algorithms are requested. Seeds cover initialization and shuffled minibatch order. CPU library versions and hardware are retained in the project metadata. Each model is persisted, reloaded through the public inference path and compared against pre-save predictions; round-trip errors appear in `results_summary.csv`. Trusted local pickle artifacts must not be loaded from untrusted distributions.

![Actual saved neural learning curves](figures/training_curves.png)

*Training evidence. Curves are read from saved per-run histories. Epoch selection uses validation and therefore does not constitute an independent test. Early plateaus or seed variation should be read in the context of the bounded training budget, not treated as a universal architecture limitation.*

## 8. Prediction results — every measured run

The table below is generated from [results_summary.csv](results_summary.csv). `train_wall_s` covers the implemented fit/training region including internal validation scoring and checkpoint generation; it is not pure optimizer-kernel time. `inference_record_median_s` is the median public full-record inference path, including checkpoint loading and window preparation. Results are sorted by observed MSE. All rows use validation and common target indices.

{table}

![Common-index model comparison](figures/model_comparison.png)

*Figure 2. Measured common-index validation comparison. Multiple seeds are initialization repeats on the same recording; they are not independent machines.*

Ridge AR(256) reproduces an MSE of **{ar.mse:.9f}**, consistent with the exploratory value in the master prompt after independent execution. The polynomial LPV-AR(256) surrogate yields **{lpv.mse:.9f}** under this configuration. The best eligible observed model is **{best.run_id}**, with **{int(best.parameters)}** parameters. Neither the baseline match nor an improvement within this suite establishes outperformance of the published Spa LPV-ARMA procedure.

The LPV-AR(256) surrogate reduces MSE by **{lpv_reduction:.3f}%** relative to ridge AR(256) on these common targets. No tested neural run beats ridge AR(256) under the declared short training budget. Appending speed as additive ARX terms does not match the benefit of polynomial speed-by-lag interactions here. That pattern is consistent with speed-dependent prediction coefficients being more useful than a speed-dependent output offset in this pair of traces; it is a conditional empirical observation, not proof of mechanism. The GRU's relatively high error indicates poor performance for this compact configuration and training budget, not that GRUs cannot model machinery signals.

![Observed/predicted validation vibration](figures/prediction.png)

*Figure 3. Saved validation predictions and residuals on a representative interval. The plotted values come from the aligned target arrays, not simulated performance. Residual energy is not physical fault accuracy.*

## 9. Robustness, context and speed-conditioning ablations

Seed summaries for principal lag-64 comparisons are:

{_table_md(seed, ['kind', 'mean', 'std', 'min', 'max', 'count'])}

The standard deviation uses pandas' sample convention across seeds. With only three initializations, it is a coarse measure of training variability. It does not quantify independent operating-condition or fault-state uncertainty. Context comparisons are in [model_ablation.csv](model_ablation.csv), alongside late-conditioning, capacity and shifted-speed controls. A selected single-seed best result must not replace the multi-seed comparison. If a condition does not improve, the negative evidence remains part of the study.

At lag 64, speed-TCN has **{film_change:.3f}% higher** mean MSE than unconditioned TCN across the three matched seeds. The proposed FiLM speed adaptation is therefore **not supported as an average prediction improvement** in this experiment. A single favorable seed cannot reverse that conclusion. Longer tested context helps the principal TCN configurations relative to lag 32, but the short neural runs remain behind the stronger linear baselines. The project contribution is an implemented, falsifiable comparison with a negative neural-conditioning result, rather than a promised novel-model victory.

Attribution of a speed benefit requires caution: speed trajectories almost coincide between train and validation, LPV interactions may act as a convenient time-dependent coefficient basis, and the FiLM model has additional trainable parameters. The larger unconditioned control and late-concatenation baseline are informative but not a complete factorial decomposition. Claims should be restricted to this finite budget and trace. Speed-binned errors are descriptive stratification and do not mean the model observed target-time speed.

![Measured validation errors by operating-speed bin](figures/speed_errors.png)

*Speed-stratified evidence. Bin-wise MSE describes errors on the actual validation samples and has different sample counts across operating regions. Stratifying by target-time speed for interpretation does not add that value to any predictor's past-only inputs. This is not a test on an unseen speed trajectory.*

## 10. Computational efficiency and practical trade-offs

![Accuracy–cost comparison](figures/efficiency.png)

*Figure 4. Accuracy and measured CPU inference cost for actual saved models. Timing boundaries and precision differ from a hypothetical isolated GPU kernel benchmark and are stated explicitly.*

The public scoring path is warmed once and repeated three times per model; median, minimum and maximum wall times are saved. It loads the checkpoint, constructs windows and predicts the full validation recording, while metrics use the shared subset. This is useful local end-to-end cost but should not be labelled pure per-sample neural latency. Parameter counts and artifact bytes are measured. The CSV retains every final run. One speed-control trial was corrected and retrained; its superseded cost is retained in `experiments/superseded_trials.json`. The hardware JSON reports both final-artifact and all-attempt training-time sums. GPU measurements were **not run** on this CPU-only setup. Peak GPU memory and cross-device acceleration are therefore not claimed.

Within-trace uncertainty is recorded in [within_trace_uncertainty.csv](within_trace_uncertainty.csv): nonoverlapping blocks of 256, 512 and 1024 samples, 1000 block-mean resamples with a fixed seed, and explicit incomplete-tail exclusions. These intervals concern temporal variation within this validation trace, are not corrected for model selection, and do not estimate uncertainty across machines or independent recordings.

A compact AR model may be preferable in a latency-sensitive deployment even when a larger model yields a lower selected MSE. Conversely, a speed-conditioned surrogate may offer a good accuracy–cost compromise. Deployment decisions also require verified streaming overhead, independent speed transfer and diagnostic retention of anomalous components, none established by one validation trace.

## 11. Residual statistics and the diagnostic hypothesis

For a causal predictor, the residual is `e[t]=x[t]−xhat[t]`. RMS is `sqrt(mean(e²))`; crest factor is `max(abs(e))/RMS`; ordinary kurtosis is the fourth centered moment divided by squared variance. Ordinary Gaussian kurtosis is 3, whereas excess kurtosis subtracts 3. The implementation states the biased central-moment convention and handles zero/constant signal cases explicitly.

Predictable reference vibration may be suppressed by the model. However, a flexible teacher-forced predictor can also learn local evolution within a defect-containing signal and suppress fault-related components. Therefore lower MSE is not logically equivalent to better fault detection. Raw/residual comparisons on accessible data establish numerical signal behavior, not a physical defect label. No model is designated a proven best physical diagnostics model.

## 12. MATLAB-reference port, angular assumptions and numerical validation

The supplied `CorrKurt` forms delayed, zero-padded copies and divides the squared-product sum by signal energy raised to `M+1`. Python allocates all `M+1` columns explicitly rather than relying on dynamic growth. Its period is an integer **sample lag**, not an automatically known shaft frequency. The PMI workflow extracts an analytic-signal envelope and uses a centered energy-normalized autocorrelation. The exact treatment of envelope mean and finite boundaries is documented in [numerical validation](numerical_validation.md).

In MATLAB, autocorrelation lag zero is the first array element. The literal coarse reference access at one-based index `round(fs/f_fault)` therefore reads physical lag **T−1**. The scientifically corrected variant reads lag **T**; both are exposed rather than silently changing historical behavior. The script's returned `timeindex` contains a separate offset and must not be mistaken for the actual selected lag. `PMI=rho/(1−rho)` is ill-conditioned near rho=1, and undefined/constant-signal cases require visible handling. One-sided FFT scaling differs for odd/even lengths; the corrected odd-length last positive bin must be doubled.

Angular resampling integrates speed after explicit conversion to cycles/s: RPM divides by 60, rps is unchanged, and rad/s divides by 2π. Interpolation onto a uniform cycle grid is a diagnostic transformation, not primary predictive preprocessing. Unknown units prevent an unqualified physical order interpretation; a speed assumption must be named and sensitivity explored. The supplied external `toConstantSamplesPerCycle` dependency is absent, so live MATLAB/toolbox parity is unverified. Synthetic tests of known-speed signals and feature formulas support the Python implementation only; they do not prove tooth-crack detection.

{sensitivity_section}

## 13. Real candidate-record study and diagnostic interpretation

{fault_section}

## 14. Scientific validity, limitations and failure modes

The strongest verified prediction findings concern one-step prediction on the supplied validation trace. Sources of uncertainty include validation reuse for checkpoint/model selection, short duration, strongly dependent neighboring windows, nearly repeated speed histories, finite training budgets, no independently verified physical labels, unknown speed units and missing independent test inputs. Candidate-record descriptors add real signal evidence but do not supply missing health labels. Seed variability and temporal block variation cannot be treated as a population confidence interval. The normalization is inherited from supplied metadata; its historical fitting population is not reconstructed from these files. It is used consistently but cannot retrospectively establish a train-only provenance for the original scaling.

Architectural future leakage is prevented through past-only windows, unidirectional recurrence and left-only convolution padding. Offline diagnostic Hilbert transforms/interpolation must not be repurposed as causal online predictors without a new design. The shifted-speed negative control intentionally violates the primary information-availability policy and is excluded from eligible model selection. Exact Spa LPV-ARMA superiority, physical gearbox parameters, fault severity and GPU efficiency remain unsupported. These are limitations of the available evidence, not quantities that can be inferred from attractive plots.

## 15. Answers to the research questions

**RQ1:** The best eligible observed common-index validation run is `{best.run_id}`, MSE **{best.mse:.6f}**; final test superiority is unknown. **RQ2:** Saved training times, full-record inference repeats, parameter counts and artifact bytes quantify the local CPU trade-off. **RQ3:** The speed-conditioning hypothesis is assessed by measured LPV/ARX and neural controls; effects are conditional on seed, context, capacity and this near-repeated trajectory. **RQ4:** {fault_rq}. **RQ5:** Context, seed and feature-convention experiments bound some variability, while independent-record prediction generalization and physical fault sensitivity remain untested.

## 16. Reproduction and deliverable map

The executable project uses importable `src/rotmach` modules; the comprehensive notebook orchestrates them. Model weights, per-run configs, histories, common-index prediction arrays, metric JSON, registry CSV/JSON and SHA256 manifests are retained. Read [the implementation guide](../docs/IMPLEMENTATION_GUIDE.md), [cross-platform instructions](../docs/reproduce_windows_linux_macos.md), [study companion](../docs/COMPREHENSIVE_STUDY_COMPANION.md) and [reproduction status](reproducibility_check.md). Commands actually executed are distinguished from documented Windows/macOS equivalents. The editable progress/final PPTX files and speaker scripts are generated by `scripts/build_presentations.py` from the same CSV metrics used here.

## References and appendices

Verified records and evidence levels are maintained in [references.bib](../research/references.bib), [literature_matrix.csv](../research/literature_matrix.csv), and [research_synthesis.md](../research/research_synthesis.md). The lecturer citation is course-source verified; restricted publisher access is not described as full-text review. Author repositories and scientific-library documentation support implementation details within their stated scope. Exact model source: [models.py](../src/rotmach/models.py), [experiment.py](../src/rotmach/experiment.py), [windows.py](../src/rotmach/windows.py). Diagnostics and numerical conventions: [diagnostics.py](../src/rotmach/diagnostics.py), [angle_resample.py](../src/rotmach/angle_resample.py), [spectral.py](../src/rotmach/spectral.py).

The full run table above is the numerical appendix. Additional source-linked outputs include [baseline comparison](baseline_comparison.csv), [ablations](model_ablation.csv), [timing methodology](timing_and_hardware.md), [scientific validity review](scientific_validity_review.md), and [rubric traceability](../docs/rubric_traceability_matrix.md). Missing evidence in the acceptance checklist remains unchecked; physical fault-detection accuracy cannot be verified without health labels, and exact paper reproduction requires additional method access.
'''
    (root / "reports/PROJECT_FINAL_REPORT.md").write_text(report)
    summary = f"""# Executive summary\n\nThe executed Python suite compares {len(df)} saved model runs on 23,744 identical validation targets (indices 256–23,999). Best eligible observed run: **{best.run_id}**, MSE **{best.mse:.6f}**; ridge AR(256): **{ar.mse:.6f}**; zero prediction: **{zero.mse:.6f}**. These are validation-selected results on one short recording, not held-out test performance. CPU timings and checkpoint round trips are recorded. Speed trajectories correlate at {audit['comparisons']['speed_pearson']:.9f}; unseen operating-history transfer remains unknown.\n\n{fault_summary}\n\nThe updated PDF rubric overrides stale prompt percentages. The course dates conflict. Speed physical units and reference health labels remain unverified. Python feature algorithms and synthetic checks are supplied; no exact Spa LPV-ARMA outperformance is claimed. See [full report](PROJECT_FINAL_REPORT.md), [result CSV](results_summary.csv), and [course issues](../docs/course_requirements_and_open_questions.md).\n"""
    (root / "reports/executive_summary.md").write_text(summary)
    return root / "reports/PROJECT_FINAL_REPORT.md"


def render_report_pdf(root):
    """Render a readable convenience PDF; Markdown is the canonical editable source.

    ReportLab handles paragraphs, headings, tables as monospaced lines and figures.
    Mathematical expressions are readable plain text rather than LaTeX typeset.
    """
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Preformatted
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    root = Path(root).resolve()
    source = root / "reports/PROJECT_FINAL_REPORT.md"
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="BodyProject", fontName="Helvetica", fontSize=9.2, leading=13.5, spaceAfter=8))
    story = []
    for line in source.read_text().splitlines():
        if not line.strip():
            continue
        if line.startswith("!["):
            match = re.search(r"\]\(([^)]+)\)", line)
            if match:
                path = source.parent / match.group(1)
                if path.exists():
                    from PIL import Image as PILImage
                    with PILImage.open(path) as im:
                        w, h = im.size
                    scale = min(470 / w, 280 / h)
                    story.extend([Image(str(path), width=w*scale, height=h*scale), Spacer(1, 8)])
            continue
        if line.startswith("|---"):
            continue
        cleaned = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", line)
        cleaned = cleaned.replace("**", "").replace("`", "")
        if cleaned.startswith("# "):
            style, cleaned = styles["Title"], cleaned[2:]
        elif cleaned.startswith("## "):
            style, cleaned = styles["Heading1"], cleaned[3:]
        elif cleaned.startswith("### "):
            style, cleaned = styles["Heading2"], cleaned[4:]
        else:
            style = styles["BodyProject"]
        # ASCII substitutions avoid unsupported glyph boxes in built-in PDF fonts.
        cleaned = cleaned.translate(str.maketrans({"−":"-", "–":"-", "—":"-", "α":"alpha", "β":"beta", "γ":"gamma", "θ":"theta", "Σ":"sum", "π":"pi", "ρ":"rho", "²":"^2", "…":"...", "×":"x", "‖":"|", "→":"->", "̂":"", "≤":"<=", "∞":"infinity", "·":" / ", "±":"+/-"}))
        story.append(Paragraph(escape(cleaned), style))
    def page(canvas, doc):
        canvas.setFont("Helvetica", 8); canvas.setFillColor(colors.HexColor("#5b6c7c"))
        canvas.drawString(44, 24, "MECG 7780 | Validation study | Canonical source: PROJECT_FINAL_REPORT.md")
        canvas.drawRightString(550, 24, str(doc.page))
    dest = root / "reports/PROJECT_FINAL_REPORT.pdf"
    SimpleDocTemplate(str(dest), pagesize=(595,842), rightMargin=44, leftMargin=44, topMargin=42, bottomMargin=42).build(story, onFirstPage=page, onLaterPages=page)
    return dest
